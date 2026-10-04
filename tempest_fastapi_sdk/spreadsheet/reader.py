"""Read ``.xlsx`` workbooks as rows — an upload, a file, a downloaded export.

The inverse of :mod:`~tempest_fastapi_sdk.spreadsheet.writer`. Each tab is
read the way :mod:`~tempest_fastapi_sdk.spreadsheet.google` reads a CSV
export: row 1 is the header, every later row becomes a ``dict`` keyed by
it, blank rows are skipped without shifting the numbering, and the
``_as`` variant validates each row into a Pydantic model and reports the
first failure by its row number in the sheet.

What changes from CSV is the cell. A CSV carries the *formatted* text
(``"R$ 1.234,56"``, ``"04/10/2026"``, depending on the sheet's locale);
an ``.xlsx`` stores the value, so a cell comes back as the Python type
``openpyxl`` reads — see :data:`XlsxCellValue`.

    from pydantic import BaseModel

    from tempest_fastapi_sdk.spreadsheet import read_xlsx_as


    class Sale(BaseModel):
        produto: str
        total: float


    def load(data: bytes) -> list[Sale]:
        return read_xlsx_as(data, Sale, sheet="Vendas")

Needs the ``[spreadsheet]`` extra (``openpyxl``). The engine is imported
at first use, so importing this module works without it; the call raises
:class:`ImportError` naming the extra.
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Callable, Iterable, Sequence
from datetime import datetime, time, timedelta
from os import PathLike
from typing import IO, TYPE_CHECKING, Any, Final, TypeVar

from pydantic import BaseModel, ValidationError

from tempest_fastapi_sdk.exceptions.base import AppException

if TYPE_CHECKING:
    from openpyxl.workbook import Workbook

ModelT = TypeVar("ModelT", bound=BaseModel)
CellT = TypeVar("CellT")

XlsxCellValue = str | int | float | bool | datetime | time | timedelta | None
"""What a cell of a read ``.xlsx`` holds.

The type ``openpyxl`` (``data_only=True``) hands back, unchanged:

* a number is ``int`` or ``float`` — whichever the file wrote. The Google
  export writes ``30.0`` and ``50`` in the same column, so declare the
  field ``float`` or ``Decimal``, not ``int``. Currency is a plain number
  (``84.39``); a percentage is the **ratio** (``0.4`` for ``40%``).
* a date is a ``datetime`` — also for a date-only cell, at midnight; a
  time-only cell is a ``time``, a duration format a ``timedelta``.
* a formula is its **cached result**; a formula whose result is the empty
  string reads as ``None``, like an empty cell.
* an error cell (``#DIV/0!``) is its text.
"""

XlsxSource = bytes | str | PathLike[str] | IO[bytes]
"""Where a workbook is read from.

``bytes`` (an upload read into memory, an HTTP body), a filesystem path,
or a binary file object — ``UploadFile.file`` included.
"""

_MISSING_DEPENDENCY: Final[str] = (
    "openpyxl is required to read .xlsx files. Install the extra: "
    'pip install "tempest-fastapi-sdk[spreadsheet]"'
)


class InvalidSpreadsheetError(AppException):
    """The source is not an ``.xlsx`` workbook.

    Raised for anything ``openpyxl`` cannot open as one: a CSV or an old
    binary ``.xls`` renamed to ``.xlsx``, an empty upload, a ZIP that is
    not an Office document. ``details["reason"]`` carries the engine's
    message for whoever debugs it.

    The status is ``422``: in the common case — a user uploading a file
    to an endpoint — the input is what is wrong.
    """

    message: str = "The file is not a valid .xlsx spreadsheet."
    code: str = "SPREADSHEET_INVALID"
    status_code: int = 422


class SheetNotFoundError(AppException):
    """The workbook has no tab with the requested name or index.

    ``details["sheet"]`` is what was asked for and ``details["available"]``
    lists the tab names the workbook does have, in tab order. Note that a
    tab name keeps its whitespace: a tab shown as ``Setembro`` may be
    called ``"Setembro "``.
    """

    message: str = "The spreadsheet has no such sheet."
    code: str = "SPREADSHEET_SHEET_NOT_FOUND"
    status_code: int = 422


class SpreadsheetRowError(AppException):
    """One row of a workbook tab did not validate against the target model.

    The same contract as
    :class:`~tempest_fastapi_sdk.spreadsheet.google.GoogleSheetRowError`
    (which is a subclass, so ``except SpreadsheetRowError`` covers both
    readers): ``details["row"]`` is the row number the spreadsheet shows —
    header = 1 — and ``details["errors"]`` is Pydantic's error list without
    the ``url`` and ``ctx`` keys. ``details["sheet"]`` names the tab.
    """

    message: str = "A row of the spreadsheet failed validation."
    code: str = "SPREADSHEET_ROW_INVALID"
    status_code: int = 422


def _is_blank(value: object) -> bool:
    """Tell whether a cell counts as empty.

    ``None`` (an empty ``.xlsx`` cell) and ``""`` (an empty CSV cell) are
    blank; ``0`` and ``False`` are values.

    Args:
        value (object): The cell value.

    Returns:
        bool: ``True`` when the cell is empty.
    """
    return value is None or value == ""


def _number_rows(
    header: Sequence[str],
    records: Iterable[Sequence[CellT]],
    pad: CellT,
) -> list[tuple[int, dict[str, CellT]]]:
    """Key the records under a header row and number them as the sheet does.

    Shared by the CSV and the ``.xlsx`` readers, so both define the edge
    cases the same way: a record blank within the header's width is
    skipped but still counts toward the numbering; a record shorter than
    the header is padded with ``pad``; cells beyond the header's width are
    ignored. Header names are taken verbatim, so an empty header cell is
    the key ``""`` and a repeated name keeps the **last** column's value.

    Args:
        header (Sequence[str]): The column names (row 1).
        records (Iterable[Sequence[CellT]]): The rows after the header, in
            order, one per spreadsheet row — blank ones included.
        pad (CellT): Value for the cells a short record lacks.

    Returns:
        list[tuple[int, dict[str, CellT]]]: ``(row_number, row)`` pairs,
        the header being row 1.
    """
    width: int = len(header)
    rows: list[tuple[int, dict[str, CellT]]] = []
    for offset, record in enumerate(records):
        cells: list[CellT] = list(record[:width])
        if all(_is_blank(cell) for cell in cells):
            continue
        cells.extend([pad] * (width - len(cells)))
        rows.append((offset + 2, dict(zip(header, cells, strict=True))))
    return rows


def _validate_rows(
    rows: Iterable[tuple[int, dict[str, CellT]]],
    schema: type[ModelT],
    *,
    omit_blank: bool,
    row_error: Callable[[int, list[dict[str, Any]]], AppException],
) -> list[ModelT]:
    """Validate numbered rows into ``schema``, stopping at the first failure.

    Args:
        rows (Iterable[tuple[int, dict[str, CellT]]]): Output of
            :func:`_number_rows`.
        schema (type[ModelT]): The model each row validates into.
        omit_blank (bool): Drop blank cells (``None`` / ``""``) before
            validating, so a blank cell means *absent*.
        row_error (Callable[[int, list[dict[str, Any]]], AppException]):
            Builds the exception for a failing row from its number and
            Pydantic's error list.

    Returns:
        list[ModelT]: One instance per row, in order.

    Raises:
        AppException: Whatever ``row_error`` builds, chained to the
            :class:`pydantic.ValidationError`.
    """
    models: list[ModelT] = []
    for row_number, row in rows:
        payload: dict[str, CellT] = (
            {key: value for key, value in row.items() if not _is_blank(value)}
            if omit_blank
            else row
        )
        try:
            models.append(schema.model_validate(payload))
        except ValidationError as exc:
            errors: list[dict[str, Any]] = [
                dict(error)
                for error in exc.errors(include_url=False, include_context=False)
            ]
            raise row_error(row_number, errors) from exc
    return models


def _header_name(value: XlsxCellValue) -> str:
    """Turn a header cell into a column name.

    Args:
        value (XlsxCellValue): The header cell.

    Returns:
        str: The text as written; ``""`` for an empty cell; ``str(value)``
        for a header typed as a number or a date.
    """
    if value is None:
        return ""
    return value if isinstance(value, str) else str(value)


def _require_openpyxl() -> None:
    """Fail early, naming the extra, when ``openpyxl`` is not installed.

    Called before any work that would be wasted without the engine — the
    Google reader calls it before downloading the workbook.

    Raises:
        ImportError: When the ``[spreadsheet]`` extra is not installed.
    """
    try:
        import openpyxl  # noqa: F401
    except ImportError as exc:
        raise ImportError(_MISSING_DEPENDENCY) from exc


def _open_workbook(source: XlsxSource) -> Workbook:
    """Open a workbook read-only, with formulas as their cached values.

    Args:
        source (XlsxSource): Bytes, a path, or a binary file object.

    Returns:
        Workbook: The workbook. The caller closes it.

    Raises:
        ImportError: When the ``[spreadsheet]`` extra is not installed.
        InvalidSpreadsheetError: If the source is not an ``.xlsx``.
        FileNotFoundError: If a path does not exist.
    """
    _require_openpyxl()
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    target: str | PathLike[str] | IO[bytes] = (
        io.BytesIO(source) if isinstance(source, bytes) else source
    )
    try:
        return load_workbook(target, read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError) as exc:
        raise InvalidSpreadsheetError(details={"reason": str(exc)}) from exc


def _sheet_rows(
    worksheet: Any,
) -> tuple[str, list[tuple[int, dict[str, XlsxCellValue]]]]:
    """Read one tab into numbered rows.

    The declared dimensions are reset first: ``openpyxl`` in read-only mode
    trusts the ``<dimension>`` tag, and a stale one truncates the rows it
    yields. The Google export writes no such tag at all.

    Args:
        worksheet (Any): An ``openpyxl`` read-only worksheet. Typed as
            ``Any`` because the engine is imported lazily.

    Returns:
        tuple[str, list[tuple[int, dict[str, XlsxCellValue]]]]: The tab's
        title and its ``(row_number, row)`` pairs.
    """
    worksheet.reset_dimensions()
    records: Iterable[Sequence[XlsxCellValue]] = worksheet.iter_rows(values_only=True)
    iterator = iter(records)
    first: Sequence[XlsxCellValue] | None = next(iterator, None)
    if first is None:
        return worksheet.title, []
    header: list[str] = [_header_name(cell) for cell in first]
    return worksheet.title, _number_rows(header, iterator, None)


def _pick_sheet(workbook: Workbook, sheet: str | int) -> Any:
    """Find a tab by name or by position.

    Args:
        workbook (Workbook): The open workbook.
        sheet (str | int): Tab name, or 0-based position in tab order.

    Returns:
        Any: The read-only worksheet.

    Raises:
        SheetNotFoundError: If no worksheet has that name or position.
    """
    worksheets: list[Any] = list(workbook.worksheets)
    if isinstance(sheet, str):
        for worksheet in worksheets:
            if worksheet.title == sheet:
                return worksheet
    elif -len(worksheets) <= sheet < len(worksheets):
        return worksheets[sheet]
    raise SheetNotFoundError(
        message=f"The spreadsheet has no sheet {sheet!r}.",
        message_params={"sheet": sheet},
        details={
            "sheet": sheet,
            "available": [worksheet.title for worksheet in worksheets],
        },
    )


def _read_sheet(
    source: XlsxSource,
    sheet: str | int,
) -> tuple[str, list[tuple[int, dict[str, XlsxCellValue]]]]:
    """Open the workbook, read one tab and close it.

    Args:
        source (XlsxSource): Bytes, a path, or a binary file object.
        sheet (str | int): Tab name or 0-based position.

    Returns:
        tuple[str, list[tuple[int, dict[str, XlsxCellValue]]]]: The tab's
        title and its numbered rows.
    """
    workbook: Workbook = _open_workbook(source)
    try:
        return _sheet_rows(_pick_sheet(workbook, sheet))
    finally:
        workbook.close()


def read_xlsx(
    source: XlsxSource,
    *,
    sheet: str | int = 0,
) -> list[dict[str, XlsxCellValue]]:
    """Read one tab of an ``.xlsx`` workbook as a list of rows.

    Row 1 is the header; each later row is a ``dict`` keyed by it, each
    value an :data:`XlsxCellValue` — an empty cell is ``None``. Blank rows
    are skipped. A tab with only a header, or nothing at all, returns
    ``[]``. Merged cells keep the value in their top-left cell only; the
    rest of the range reads ``None``.

    Args:
        source (XlsxSource): The workbook as bytes, a path, or a binary
            file object.
        sheet (str | int): The tab's name, or its 0-based position in tab
            order. Defaults to the first tab.

    Returns:
        list[dict[str, XlsxCellValue]]: The data rows, in sheet order.

    Raises:
        ImportError: When the ``[spreadsheet]`` extra is not installed.
        InvalidSpreadsheetError: If the source is not an ``.xlsx``.
        SheetNotFoundError: If the workbook has no such tab.
        FileNotFoundError: If a path does not exist.
    """
    _, rows = _read_sheet(source, sheet)
    return [row for _, row in rows]


def read_xlsx_sheets(
    source: XlsxSource,
) -> dict[str, list[dict[str, XlsxCellValue]]]:
    """Read every tab of an ``.xlsx`` workbook, keyed by tab name.

    Each tab follows the rules of :func:`read_xlsx`. Chart sheets, which
    hold no cells, are left out.

    Args:
        source (XlsxSource): The workbook as bytes, a path, or a binary
            file object.

    Returns:
        dict[str, list[dict[str, XlsxCellValue]]]: The rows of each tab, in
        tab order. An empty tab maps to ``[]``.

    Raises:
        ImportError: When the ``[spreadsheet]`` extra is not installed.
        InvalidSpreadsheetError: If the source is not an ``.xlsx``.
        FileNotFoundError: If a path does not exist.
    """
    workbook: Workbook = _open_workbook(source)
    try:
        sheets: dict[str, list[dict[str, XlsxCellValue]]] = {}
        for worksheet in workbook.worksheets:
            title, rows = _sheet_rows(worksheet)
            sheets[title] = [row for _, row in rows]
        return sheets
    finally:
        workbook.close()


def read_xlsx_as(
    source: XlsxSource,
    schema: type[ModelT],
    *,
    sheet: str | int = 0,
    omit_blank: bool = True,
) -> list[ModelT]:
    """Read one tab of an ``.xlsx`` workbook and validate each row.

    The header names are the keys handed to ``schema.model_validate``, so
    a column maps to the field of the same name (use
    ``validation_alias`` for a header that is not a valid identifier).
    Cells arrive typed (:data:`XlsxCellValue`), so a currency column
    validates into ``float`` or ``Decimal`` without undoing any formatting.

    Args:
        source (XlsxSource): The workbook as bytes, a path, or a binary
            file object.
        schema (type[ModelT]): The Pydantic model each row validates into.
        sheet (str | int): The tab's name, or its 0-based position.
            Defaults to the first tab.
        omit_blank (bool): Drop empty cells from the row before
            validating, so a blank cell means *absent*: the field falls
            back to its default, and a required field reports ``missing``.
            ``False`` passes ``None`` through.

    Returns:
        list[ModelT]: One instance per data row, in sheet order. ``[]``
        when the tab has no data row.

    Raises:
        ImportError: When the ``[spreadsheet]`` extra is not installed.
        InvalidSpreadsheetError: If the source is not an ``.xlsx``.
        SheetNotFoundError: If the workbook has no such tab.
        SpreadsheetRowError: On the first row that fails validation;
            ``details["row"]`` is its number in the sheet (header = 1) and
            ``details["sheet"]`` the tab's name.
        FileNotFoundError: If a path does not exist.
    """
    title, rows = _read_sheet(source, sheet)

    def row_error(row_number: int, errors: list[dict[str, Any]]) -> AppException:
        """Build the error for a failing row of this tab.

        Args:
            row_number (int): The row's number in the sheet.
            errors (list[dict[str, Any]]): Pydantic's errors for the row.

        Returns:
            AppException: The :class:`SpreadsheetRowError`.
        """
        return SpreadsheetRowError(
            message=f"Row {row_number} of sheet {title!r} failed validation.",
            message_params={"row": row_number, "sheet": title},
            details={"row": row_number, "sheet": title, "errors": errors},
        )

    return _validate_rows(rows, schema, omit_blank=omit_blank, row_error=row_error)


__all__: list[str] = [
    "InvalidSpreadsheetError",
    "SheetNotFoundError",
    "SpreadsheetRowError",
    "XlsxCellValue",
    "XlsxSource",
    "read_xlsx",
    "read_xlsx_as",
    "read_xlsx_sheets",
]
