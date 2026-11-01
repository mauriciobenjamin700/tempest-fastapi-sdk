"""Write a deck as a Word handout, and an answer key as a Word document.

The handout puts one slide per page with room for the presenter: the
slide's content, its speaker notes (or ruled lines to write on when there
are none) and a footer with the slide number. Ported from
``slide_docx_service.py`` and ``answer_key_docx_service.py`` of Leviathan's
``slide-generator``.

``python-docx`` is imported inside each function, so importing this module
does not need the ``[slides]`` extra.
"""

from __future__ import annotations

import io
from typing import TYPE_CHECKING

from tempest_fastapi_sdk.slides.exceptions import missing_extra_message
from tempest_fastapi_sdk.slides.html_renderer import strip_tags
from tempest_fastapi_sdk.slides.schemas import (
    AnswerKey,
    Slide,
    SlideDeck,
    SlideKindEnum,
)
from tempest_fastapi_sdk.slides.vocabulary import AnswerKeyVocabulary, SlideVocabulary

if TYPE_CHECKING:
    from docx.document import Document as DocxDocument
    from docx.text.run import Run

TITLE_COLOR: str = "1F2937"
"""Slide titles in the handout."""

ACCENT_COLOR: str = "2563EB"
"""Eyebrows, notes heading, answer-key sections and answer labels."""

MUTED_COLOR: str = "6B7280"
"""Cover subtitle, placeholders and footers."""

NOTES_COLOR: str = "334155"
"""Speaker notes in the handout and answers in the answer key."""

ANSWER_KEY_TITLE_COLOR: str = "1E3A8A"
"""Title of the answer-key document."""

QUESTION_COLOR: str = "0F172A"
"""Questions of the answer-key document."""

BLANK_NOTE_LINES: int = 4
"""Ruled lines printed in place of absent speaker notes."""


def _style(
    run: Run,
    *,
    size: float,
    color: str | None = None,
    bold: bool = False,
    italic: bool = False,
) -> None:
    """Apply size, colour and weight to a run.

    Args:
        run (Run): The python-docx run.
        size (float): Font size, in points.
        color (str | None): ``RRGGBB`` colour; ``None`` keeps the default.
        bold (bool): Whether the run is bold.
        italic (bool): Whether the run is italic.
    """
    from docx.shared import Pt, RGBColor

    run.font.size = Pt(size)
    run.font.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def _page_break(document: DocxDocument) -> None:
    """Start a new page.

    Args:
        document (DocxDocument): The document.
    """
    from docx.enum.text import WD_BREAK

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _eyebrow(slide: Slide, vocabulary: SlideVocabulary) -> str:
    """Return the label printed above a slide's title.

    Args:
        slide (Slide): The slide.
        vocabulary (SlideVocabulary): Labels in use.

    Returns:
        str: The label, uppercased, or ``""`` when none applies.
    """
    if slide.kind == SlideKindEnum.LESSON_TITLE and slide.subtitle:
        return slide.subtitle.upper()
    if slide.kind == SlideKindEnum.CONCLUSION:
        return vocabulary.conclusion_label.upper()
    if slide.lesson_index is not None:
        return vocabulary.lesson(slide.lesson_index).upper()
    return ""


def _slide_page(
    document: DocxDocument, slide: Slide, deck: SlideDeck, vocabulary: SlideVocabulary
) -> None:
    """Write one slide as one handout page.

    Args:
        document (DocxDocument): The document.
        slide (Slide): The slide.
        deck (SlideDeck): The deck, for the footer.
        vocabulary (SlideVocabulary): Labels in use.
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    eyebrow = _eyebrow(slide, vocabulary)
    if eyebrow:
        _style(
            document.add_paragraph().add_run(eyebrow),
            size=11,
            color=ACCENT_COLOR,
            bold=True,
        )
    _style(
        document.add_paragraph().add_run(slide.title),
        size=22,
        color=TITLE_COLOR,
        bold=True,
    )
    if slide.kind == SlideKindEnum.LESSON_TITLE:
        run = document.add_paragraph().add_run(vocabulary.handout_cover_note)
        _style(run, size=11, color=MUTED_COLOR, italic=True)
    elif slide.kind == SlideKindEnum.BULLETS:
        for bullet in slide.bullets:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if bullet.term:
                _style(paragraph.add_run(f"{bullet.term}: "), size=12, bold=True)
            _style(paragraph.add_run(bullet.description), size=12)
    else:
        for text in slide.paragraphs:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.space_after = Pt(8)
            _style(paragraph.add_run(text), size=12)
    document.add_paragraph()
    heading = document.add_paragraph().add_run(vocabulary.notes_heading)
    _style(heading, size=11, color=ACCENT_COLOR, bold=True)
    lines = [line.strip() for line in (slide.notes or "").splitlines() if line.strip()]
    if lines:
        for line in lines:
            _style(document.add_paragraph().add_run(line), size=11, color=NOTES_COLOR)
    else:
        for _ in range(BLANK_NOTE_LINES):
            run = document.add_paragraph().add_run("_" * 32)
            _style(run, size=11, color=MUTED_COLOR)
    footer = document.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label = (
        f"{deck.title}  ·  {vocabulary.slide_label} {slide.slide_number} / {deck.total}"
    )
    _style(footer.add_run(label), size=9, color=MUTED_COLOR, italic=True)


def write_slides_docx(deck: SlideDeck, *, vocabulary: SlideVocabulary) -> bytes:
    """Write a deck as a presenter handout, one slide per page.

    Args:
        deck (SlideDeck): The deck.
        vocabulary (SlideVocabulary): Labels to print.

    Returns:
        bytes: The ``.docx`` file.

    Raises:
        ImportError: When ``python-docx`` is missing.
    """
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - extra-gated
        raise ImportError(missing_extra_message("python-docx")) from exc
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    document = Document()
    cover = document.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style(cover.add_run(deck.title), size=28, color=TITLE_COLOR, bold=True)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style(
        subtitle.add_run(vocabulary.handout_subtitle),
        size=12,
        color=MUTED_COLOR,
        italic=True,
    )
    for slide in deck.slides:
        _page_break(document)
        _slide_page(document, slide, deck, vocabulary)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def write_answer_key_docx(
    answer_key: AnswerKey, *, vocabulary: AnswerKeyVocabulary
) -> bytes:
    """Write an answer key as a Word document.

    Args:
        answer_key (AnswerKey): The answer key.
        vocabulary (AnswerKeyVocabulary): Labels to print.

    Returns:
        bytes: The ``.docx`` file.

    Raises:
        ImportError: When ``python-docx`` is missing.
    """
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - extra-gated
        raise ImportError(missing_extra_message("python-docx")) from exc
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    document = Document()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style(
        title.add_run(answer_key.title),
        size=24,
        color=ANSWER_KEY_TITLE_COLOR,
        bold=True,
    )
    if answer_key.subtitle:
        subtitle = document.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _style(
            subtitle.add_run(answer_key.subtitle),
            size=11,
            color=MUTED_COLOR,
            italic=True,
        )
    for section in answer_key.sections:
        document.add_paragraph()
        if section.title:
            _style(
                document.add_paragraph().add_run(section.title),
                size=14,
                color=ACCENT_COLOR,
                bold=True,
            )
        for item in section.items:
            question = document.add_paragraph()
            question.paragraph_format.space_before = Pt(8)
            _style(
                question.add_run(strip_tags(item.question)),
                size=12,
                color=QUESTION_COLOR,
                bold=True,
            )
            answer = document.add_paragraph()
            answer.paragraph_format.left_indent = Pt(18)
            _style(
                answer.add_run(f"{vocabulary.answer_label} "),
                size=12,
                color=ACCENT_COLOR,
                bold=True,
            )
            _style(
                answer.add_run(strip_tags(item.answer)),
                size=12,
                color=NOTES_COLOR,
                italic=True,
            )
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
