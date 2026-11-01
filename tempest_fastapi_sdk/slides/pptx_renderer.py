"""Write a :class:`SlideDeck` as a PowerPoint file.

Every slide is drawn on the blank layout with explicit text boxes, so the
result looks the same in PowerPoint, LibreOffice Impress and Google Slides
instead of depending on each one's default placeholders. Speaker notes go
to the notes pane. Ported from ``slide_pptx_service.py`` of Leviathan's
``slide-generator``; geometry, sizes and colours are the original ones.

``python-pptx`` is imported inside the function, so importing this module
does not need the ``[slides]`` extra.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from typing import TYPE_CHECKING

from tempest_fastapi_sdk.slides.exceptions import missing_extra_message
from tempest_fastapi_sdk.slides.schemas import (
    Slide,
    SlideDeck,
    SlideKindEnum,
    SlideThemeEnum,
)
from tempest_fastapi_sdk.slides.vocabulary import SlideVocabulary

if TYPE_CHECKING:
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_PARAGRAPH_ALIGNMENT
    from pptx.slide import Slide as PptxSlide
    from pptx.text.text import TextFrame, _Paragraph

BLANK_LAYOUT_INDEX: int = 6
"""Index of the blank layout in python-pptx's default template."""

SLIDE_WIDTH_INCHES: float = 13.333
"""16:9 slide width."""

SLIDE_HEIGHT_INCHES: float = 7.5
"""16:9 slide height."""

MARGIN_INCHES: float = 0.7
"""Left and right margin of every text box."""

FONT_NAME: str = "Calibri"
"""Font of every run; present in Office and metric-compatible elsewhere."""


@dataclass(frozen=True, slots=True)
class PptxPalette:
    """Colours of one theme, as ``RRGGBB`` hex strings.

    Attributes:
        background (str): Slide background.
        text (str): Body and title text.
        muted (str): Footer text.
        accent (str): Eyebrows, bullet markers and terms.
        hero_background (str): Background of a lesson-title slide.
        hero_text (str): Text of a lesson-title slide.
    """

    background: str
    text: str
    muted: str
    accent: str
    hero_background: str
    hero_text: str


PPTX_PALETTES: dict[SlideThemeEnum, PptxPalette] = {
    SlideThemeEnum.DEFAULT: PptxPalette(
        background="FFFFFF",
        text="1F2937",
        muted="6B7280",
        accent="2563EB",
        hero_background="1E3A8A",
        hero_text="FFFFFF",
    ),
    SlideThemeEnum.DARK: PptxPalette(
        background="0F172A",
        text="E2E8F0",
        muted="94A3B8",
        accent="60A5FA",
        hero_background="020617",
        hero_text="F8FAFC",
    ),
}
"""Palette per theme. Ported from ``_THEMES`` in Leviathan's
``slide_pptx_service.py``."""


class _Writer:
    """Draws slides into one presentation.

    Attributes:
        deck (SlideDeck): The deck being written.
        palette (PptxPalette): Colours in use.
        vocabulary (SlideVocabulary): Labels in use.
    """

    def __init__(
        self, deck: SlideDeck, palette: PptxPalette, vocabulary: SlideVocabulary
    ) -> None:
        """Bind the writer to one deck.

        Args:
            deck (SlideDeck): The deck being written.
            palette (PptxPalette): Colours in use.
            vocabulary (SlideVocabulary): Labels in use.
        """
        self.deck: SlideDeck = deck
        self.palette: PptxPalette = palette
        self.vocabulary: SlideVocabulary = vocabulary

    def write(self) -> bytes:
        """Build the presentation in memory.

        Returns:
            bytes: The ``.pptx`` file.
        """
        try:
            from pptx import Presentation
        except ImportError as exc:  # pragma: no cover - extra-gated
            raise ImportError(missing_extra_message("python-pptx")) from exc
        from pptx.util import Inches

        presentation = Presentation()
        presentation.slide_width = Inches(SLIDE_WIDTH_INCHES)
        presentation.slide_height = Inches(SLIDE_HEIGHT_INCHES)
        layout = presentation.slide_layouts[BLANK_LAYOUT_INDEX]
        for slide in self.deck.slides:
            target = presentation.slides.add_slide(layout)
            self._draw(target, slide)
        buffer = io.BytesIO()
        presentation.save(buffer)
        return buffer.getvalue()

    def _draw(self, target: PptxSlide, slide: Slide) -> None:
        """Draw one slide according to its kind.

        Args:
            target (PptxSlide): The python-pptx slide.
            slide (Slide): The slide to draw.
        """
        is_hero = slide.kind == SlideKindEnum.LESSON_TITLE
        fill = target.background.fill
        fill.solid()
        fill.fore_color.rgb = self._rgb(
            self.palette.hero_background if is_hero else self.palette.background
        )
        if is_hero:
            self._draw_hero(target, slide)
        else:
            eyebrow = self._eyebrow(slide)
            if eyebrow:
                self._text(
                    target,
                    eyebrow,
                    top=0.5,
                    height=0.4,
                    size=14,
                    bold=True,
                    color=self.palette.accent,
                )
            self._text(
                target,
                slide.title,
                top=0.95,
                height=1.0,
                size=32,
                bold=True,
                color=self.palette.text,
            )
            if slide.kind == SlideKindEnum.BULLETS:
                self._draw_bullets(target, slide)
            else:
                self._draw_paragraphs(target, slide)
            self._draw_footer(target, slide)
        notes_frame = target.notes_slide.notes_text_frame if slide.notes else None
        if notes_frame is not None and slide.notes:
            notes_frame.text = slide.notes

    def _eyebrow(self, slide: Slide) -> str:
        """Return the label above a non-hero slide's title.

        Args:
            slide (Slide): The slide.

        Returns:
            str: The conclusion label, the lesson label, or ``""``.
        """
        if slide.kind == SlideKindEnum.CONCLUSION:
            return self.vocabulary.conclusion_label.upper()
        if slide.lesson_index is not None:
            return self.vocabulary.lesson(slide.lesson_index).upper()
        return ""

    def _draw_hero(self, target: PptxSlide, slide: Slide) -> None:
        """Draw a lesson-title slide: centred eyebrow and large title.

        Args:
            target (PptxSlide): The python-pptx slide.
            slide (Slide): The slide.
        """
        from pptx.enum.text import PP_ALIGN

        if slide.subtitle:
            self._text(
                target,
                slide.subtitle.upper(),
                top=2.7,
                height=0.5,
                size=20,
                bold=False,
                color=self.palette.hero_text,
                align=PP_ALIGN.CENTER,
            )
        self._text(
            target,
            slide.title,
            top=3.3,
            height=2.0,
            size=48,
            bold=True,
            color=self.palette.hero_text,
            align=PP_ALIGN.CENTER,
        )

    def _draw_paragraphs(self, target: PptxSlide, slide: Slide) -> None:
        """Draw the justified body paragraphs of a content slide.

        Args:
            target (PptxSlide): The python-pptx slide.
            slide (Slide): The slide.
        """
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Pt

        frame = self._frame(target, top=2.1, height=4.4)
        for index, text in enumerate(slide.paragraphs):
            paragraph = self._paragraph(frame, index)
            paragraph.alignment = PP_ALIGN.JUSTIFY
            paragraph.space_after = Pt(12)
            self._run(paragraph, text, size=20, color=self.palette.text)

    def _draw_bullets(self, target: PptxSlide, slide: Slide) -> None:
        """Draw the items of a bullets slide, term highlighted.

        Args:
            target (PptxSlide): The python-pptx slide.
            slide (Slide): The slide.
        """
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Pt

        frame = self._frame(target, top=2.1, height=4.4)
        for index, bullet in enumerate(slide.bullets):
            paragraph = self._paragraph(frame, index)
            paragraph.alignment = PP_ALIGN.JUSTIFY
            paragraph.space_after = Pt(10)
            self._run(paragraph, "• ", size=20, color=self.palette.accent, bold=True)
            if bullet.term:
                self._run(
                    paragraph,
                    f"{bullet.term}: ",
                    size=20,
                    color=self.palette.accent,
                    bold=True,
                )
            self._run(paragraph, bullet.description, size=20, color=self.palette.text)

    def _draw_footer(self, target: PptxSlide, slide: Slide) -> None:
        """Draw the deck title and ``n / total`` at the bottom.

        Args:
            target (PptxSlide): The python-pptx slide.
            slide (Slide): The slide.
        """
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches

        box = target.shapes.add_textbox(
            Inches(MARGIN_INCHES),
            Inches(SLIDE_HEIGHT_INCHES - 0.5),
            Inches(SLIDE_WIDTH_INCHES - 2 * MARGIN_INCHES),
            Inches(0.3),
        )
        paragraph = box.text_frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.LEFT
        self._run(paragraph, self.deck.title, size=10, color=self.palette.muted)
        paragraph.add_run().text = "    "
        self._run(
            paragraph,
            f"{slide.slide_number} / {self.deck.total}",
            size=10,
            color=self.palette.muted,
        )

    def _frame(self, target: PptxSlide, *, top: float, height: float) -> TextFrame:
        """Add a full-width text box and return its empty frame.

        Args:
            target (PptxSlide): The python-pptx slide.
            top (float): Distance from the top, in inches.
            height (float): Box height, in inches.

        Returns:
            TextFrame: The frame, word-wrapped and cleared to one empty
            paragraph.
        """
        from pptx.util import Inches

        box = target.shapes.add_textbox(
            Inches(MARGIN_INCHES),
            Inches(top),
            Inches(SLIDE_WIDTH_INCHES - 2 * MARGIN_INCHES),
            Inches(height),
        )
        frame: TextFrame = box.text_frame
        frame.word_wrap = True
        frame.clear()  # type: ignore[no-untyped-call]
        return frame

    @staticmethod
    def _paragraph(frame: TextFrame, index: int) -> _Paragraph:
        """Return the frame's first paragraph, or append a new one.

        A cleared frame keeps one empty paragraph, so the first item reuses
        it instead of leaving a blank line at the top of the box.

        Args:
            frame (TextFrame): The text frame.
            index (int): Position of the item being written.

        Returns:
            _Paragraph: The paragraph to write into.
        """
        if index == 0:
            return frame.paragraphs[0]
        paragraph: _Paragraph = frame.add_paragraph()  # type: ignore[no-untyped-call]
        return paragraph

    def _text(
        self,
        target: PptxSlide,
        content: str,
        *,
        top: float,
        height: float,
        size: int,
        bold: bool,
        color: str,
        align: PP_PARAGRAPH_ALIGNMENT | None = None,
    ) -> None:
        """Add a full-width text box holding one styled run.

        Args:
            target (PptxSlide): The python-pptx slide.
            content (str): The text.
            top (float): Distance from the top, in inches.
            height (float): Box height, in inches.
            size (int): Font size, in points.
            bold (bool): Whether the run is bold.
            color (str): ``RRGGBB`` colour.
            align (PP_PARAGRAPH_ALIGNMENT | None): Paragraph alignment;
                ``None`` aligns left.
        """
        from pptx.enum.text import PP_ALIGN

        frame = self._frame(target, top=top, height=height)
        paragraph = frame.paragraphs[0]
        paragraph.alignment = align if align is not None else PP_ALIGN.LEFT
        self._run(paragraph, content, size=size, color=color, bold=bold)

    def _run(
        self,
        paragraph: _Paragraph,
        text: str,
        *,
        size: int,
        color: str,
        bold: bool = False,
    ) -> None:
        """Append one styled run to a paragraph.

        Args:
            paragraph (_Paragraph): The python-pptx paragraph.
            text (str): The run's text.
            size (int): Font size, in points.
            color (str): ``RRGGBB`` colour.
            bold (bool): Whether the run is bold.
        """
        from pptx.util import Pt

        run = paragraph.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = FONT_NAME
        run.font.color.rgb = self._rgb(color)

    @staticmethod
    def _rgb(hex_color: str) -> RGBColor:
        """Convert ``RRGGBB`` to python-pptx's colour type.

        Args:
            hex_color (str): The colour.

        Returns:
            RGBColor: The same colour.
        """
        from pptx.dml.color import RGBColor

        red, green, blue = bytes.fromhex(hex_color)
        return RGBColor(red, green, blue)  # type: ignore[no-untyped-call]


def write_slides_pptx(
    deck: SlideDeck, *, theme: SlideThemeEnum, vocabulary: SlideVocabulary
) -> bytes:
    """Write a deck as a ``.pptx`` file, synchronously.

    Args:
        deck (SlideDeck): The deck.
        theme (SlideThemeEnum): Which palette to use.
        vocabulary (SlideVocabulary): Labels to print.

    Returns:
        bytes: The file.

    Raises:
        ImportError: When ``python-pptx`` is missing.
    """
    return _Writer(deck, PPTX_PALETTES[SlideThemeEnum(theme)], vocabulary).write()
