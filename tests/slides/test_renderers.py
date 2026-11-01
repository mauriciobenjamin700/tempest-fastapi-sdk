"""Tests for ``SlideRenderer`` and ``AnswerKeyRenderer``.

Every binary output is opened back with the library that reads it —
python-pptx, python-docx, pypdf — so the assertions are about what a
reader sees, not about bytes that merely start with the right magic.
"""

from __future__ import annotations

import io
import zipfile

import docx
import pptx
import pytest
from pypdf import PdfReader

from tempest_fastapi_sdk.slides import (
    AnswerKey,
    AnswerKeyItem,
    AnswerKeyRenderer,
    AnswerKeySection,
    AnswerKeyVocabulary,
    SlideDeck,
    SlideRenderer,
    SlideThemeEnum,
    SlideVocabulary,
    parse_slides,
)

DECK_SOURCE: str = (
    "Lesson 1: Logic gates\n"
    "What a gate is\n"
    "A logic gate maps binary inputs to a single binary output.\n"
    "Notes: Start from the LED example.\n"
    "\n"
    "Kinds\n"
    "NOT gate: inverts its only input, every single time.\n"
    "Note: Show the truth table.\n"
    "\n"
    "Unit conclusion\n"
    "Every digital circuit is built from a handful of gates.\n"
)

PT_LABELS: SlideVocabulary = SlideVocabulary(
    lesson_label="Aula {index}",
    conclusion_label="Conclusão",
    handout_subtitle="Material de apoio do apresentador",
    notes_heading="Notas do apresentador",
    language="pt-BR",
)


@pytest.fixture
def deck() -> SlideDeck:
    """Parse the shared English source.

    Returns:
        SlideDeck: Four slides — cover, content, bullets, conclusion.
    """
    return parse_slides(DECK_SOURCE)


@pytest.fixture
def answer_key() -> AnswerKey:
    """Build an answer key carrying inline HTML in its texts.

    Returns:
        AnswerKey: One section, one item.
    """
    return AnswerKey(
        title="Gates quiz",
        subtitle="Class 3B",
        sections=[
            AnswerKeySection(
                title="Part 1: Gates",
                items=[
                    AnswerKeyItem(
                        question="1. What does <b>NOT</b> do?",
                        answer='Inverts <span class="x">the input</span>.',
                    )
                ],
            )
        ],
    )


class TestSlideHtml:
    def test_values_are_escaped(self) -> None:
        """A script in the source cannot run in the page that embeds the preview."""
        deck = parse_slides(
            "Lesson 1: <script>alert(1)</script>\n"
            "A section\n"
            "Text <img src=x onerror=y>.\n"
        )
        html = SlideRenderer().html(deck)
        assert "<script>alert" not in html
        assert "<img src=x" not in html
        assert "&lt;script&gt;" in html

    def test_labels_and_language_come_from_the_vocabulary(
        self, deck: SlideDeck
    ) -> None:
        """Eyebrows and ``<html lang>`` follow the vocabulary."""
        html = SlideRenderer(vocabulary=PT_LABELS).html(deck)
        assert '<html lang="pt-BR">' in html
        assert "Aula 1" in html
        assert "Conclusão" in html

    def test_theme_selects_the_body_class(self, deck: SlideDeck) -> None:
        """``dark`` switches the CSS variables block."""
        assert 'class="theme-dark"' in SlideRenderer().html(
            deck, theme=SlideThemeEnum.DARK
        )


class TestSlideMarkdown:
    def test_revealjs_separators_and_notes(self, deck: SlideDeck) -> None:
        """``---`` between slides, notes under ``Note:``."""
        output = SlideRenderer().markdown(deck)
        assert output.count("\n---\n") == deck.total - 1
        assert "Note:\nStart from the LED example." in output
        assert "## Unit conclusion" in output
        assert "<small>Conclusion</small>" in output


class TestSlidePptx:
    async def test_one_slide_per_deck_slide_with_notes(self, deck: SlideDeck) -> None:
        """Speaker notes land in the notes pane."""
        payload = await SlideRenderer().pptx(deck)
        presentation = pptx.Presentation(io.BytesIO(payload))
        assert len(presentation.slides) == deck.total
        notes = [
            s.notes_slide.notes_text_frame.text
            for s in presentation.slides
            if s.has_notes_slide and s.notes_slide.notes_text_frame is not None
        ]
        assert "Start from the LED example." in notes
        assert "Show the truth table." in notes

    async def test_widescreen_geometry(self, deck: SlideDeck) -> None:
        """16:9 at 13.33 x 7.5 inches."""
        presentation = pptx.Presentation(io.BytesIO(await SlideRenderer().pptx(deck)))
        assert presentation.slide_width is not None
        assert presentation.slide_height is not None
        assert round(presentation.slide_width.inches, 2) == 13.33
        assert round(presentation.slide_height.inches, 2) == 7.5

    async def test_eyebrow_uses_the_vocabulary(self, deck: SlideDeck) -> None:
        """The conclusion eyebrow is the vocabulary's, uppercased."""
        payload = await SlideRenderer(vocabulary=PT_LABELS).pptx(
            deck, theme=SlideThemeEnum.DARK
        )
        presentation = pptx.Presentation(io.BytesIO(payload))
        texts = {
            shape.text_frame.text
            for slide in presentation.slides
            for shape in slide.shapes
            if shape.has_text_frame
        }
        assert "CONCLUSÃO" in texts
        assert "AULA 1" in texts


class TestSlideDocx:
    async def test_handout_has_a_notes_block_per_slide(self, deck: SlideDeck) -> None:
        """One notes heading per slide, absent notes replaced by ruled lines."""
        payload = await SlideRenderer(vocabulary=PT_LABELS).docx(deck)
        document = docx.Document(io.BytesIO(payload))
        texts = [p.text for p in document.paragraphs]
        assert texts.count("Notas do apresentador") == deck.total
        assert "Material de apoio do apresentador" in texts
        assert "_" * 32 in texts


class TestSlidePdf:
    async def test_one_page_per_slide(self, deck: SlideDeck) -> None:
        """The 1280x720 page box holds exactly one slide."""
        reader = PdfReader(io.BytesIO(await SlideRenderer().pdf(deck)))
        assert len(reader.pages) == deck.total
        assert "LOGIC GATES" in reader.pages[0].extract_text()


class TestAnswerKeyRenderer:
    def test_html_strips_tags_then_escapes(self, answer_key: AnswerKey) -> None:
        """Inline markup is removed, not executed and not shown."""
        html = AnswerKeyRenderer().html(answer_key)
        assert "<b>NOT</b>" not in html
        assert "&lt;b&gt;" not in html
        assert "What does NOT do?" in html

    def test_footer_is_printed_only_when_set(self, answer_key: AnswerKey) -> None:
        """``footer=None`` (the default) prints no closing line."""
        assert '<div class="document__footer">' not in AnswerKeyRenderer().html(
            answer_key
        )
        footer = AnswerKeyVocabulary(footer="Documento de apoio.")
        assert "Documento de apoio." in AnswerKeyRenderer(vocabulary=footer).html(
            answer_key
        )

    def test_markdown_quotes_the_answer(self, answer_key: AnswerKey) -> None:
        """Answers are blockquotes led by the label."""
        output = AnswerKeyRenderer(
            vocabulary=AnswerKeyVocabulary(answer_label="Resposta:")
        ).markdown(answer_key)
        assert output.startswith("# Gates quiz")
        assert "_Class 3B_" in output
        assert "> Resposta: Inverts the input." in output
        assert "<span" not in output

    async def test_docx_strips_tags(self, answer_key: AnswerKey) -> None:
        """Word cannot render HTML, so the visible text is kept."""
        document = docx.Document(io.BytesIO(await AnswerKeyRenderer().docx(answer_key)))
        text = "\n".join(p.text for p in document.paragraphs)
        assert "Answer: Inverts the input." in text
        assert "<span" not in text

    async def test_pdf_renders(self, answer_key: AnswerKey) -> None:
        """A4 PDF holding the question."""
        reader = PdfReader(io.BytesIO(await AnswerKeyRenderer().pdf(answer_key)))
        assert "What does NOT do?" in reader.pages[0].extract_text()


class TestRendererConfiguration:
    @pytest.mark.parametrize("factory", [SlideRenderer, AnswerKeyRenderer])
    def test_concurrency_below_one_is_refused(self, factory: type) -> None:
        """A zero semaphore would deadlock every render."""
        with pytest.raises(ValueError, match="max_concurrent"):
            factory(max_concurrent=0)

    async def test_pptx_is_a_valid_office_package(self, deck: SlideDeck) -> None:
        """The zip carries the parts every Office reader opens first."""
        names = zipfile.ZipFile(io.BytesIO(await SlideRenderer().pptx(deck))).namelist()
        assert "[Content_Types].xml" in names
        assert "ppt/presentation.xml" in names


class TestMissingExtra:
    def test_missing_engine_names_the_extra(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """Without python-pptx the error says which extra to install."""
        import builtins
        from typing import Any

        real_import = builtins.__import__

        def refuse_pptx(name: str, *args: Any, **kwargs: Any) -> Any:
            """Fail the python-pptx import only.

            Args:
                name (str): Module name.
                *args (Any): Forwarded.
                **kwargs (Any): Forwarded.

            Returns:
                Any: The imported module.

            Raises:
                ImportError: For ``pptx``.
            """
            if name == "pptx" or name.startswith("pptx."):
                raise ImportError(name)
            return real_import(name, *args, **kwargs)

        monkeypatch.setattr(builtins, "__import__", refuse_pptx)
        from tempest_fastapi_sdk.slides.pptx_renderer import write_slides_pptx
        from tempest_fastapi_sdk.slides.vocabulary import DEFAULT_SLIDE_VOCABULARY

        with pytest.raises(ImportError, match=r"tempest-fastapi-sdk\[slides\]"):
            write_slides_pptx(
                parse_slides(DECK_SOURCE),
                theme=SlideThemeEnum.DEFAULT,
                vocabulary=DEFAULT_SLIDE_VOCABULARY,
            )


class TestUntitledSection:
    async def test_no_empty_heading_for_questions_without_a_part(self) -> None:
        """A key whose questions sit under no heading writes no blank title."""
        key = AnswerKey(
            title="Quiz",
            sections=[
                AnswerKeySection(
                    title="",
                    items=[AnswerKeyItem(question="1. Why?", answer="Because.")],
                )
            ],
        )
        renderer = AnswerKeyRenderer()
        assert "\n## \n" not in renderer.markdown(key)
        document = docx.Document(io.BytesIO(await renderer.docx(key)))
        bold_empty = [
            p for p in document.paragraphs if not p.text and any(r.bold for r in p.runs)
        ]
        assert bold_empty == []
